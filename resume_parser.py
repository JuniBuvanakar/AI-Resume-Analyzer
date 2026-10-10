from io import BytesIO
from pypdf import PdfReader
from docx import Document

def extract_resume_text(uploaded_file, filename: str) -> str:
    """Extract text from a PDF or DOCX Streamlit UploadedFile."""
    data = uploaded_file.getvalue()
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("Only PDF and DOCX files are supported.")
